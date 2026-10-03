from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document('Downloads/Proposal_Penelitian_Tia_Bab1.docx')

def add_after(doc, ref_idx, text, bold=False, italic=False, align='left', size_pt=11, space_after=6, indent_first=True):
    ref_p = doc.paragraphs[ref_idx]._element
    new_p = doc.add_paragraph()
    ref_p.getparent().insert(ref_p.getparent().index(ref_p) + 1, new_p._element)
    new_p.clear()
    run = new_p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size_pt)
    if bold: run.bold = True
    if italic: run.italic = True
    run.font.color.rgb = RGBColor(0,0,0)
    if align == 'center': new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'justify': new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else: new_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    new_p.paragraph_format.space_after = Pt(space_after)
    new_p.paragraph_format.line_spacing = 1.5
    if indent_first: new_p.paragraph_format.first_line_indent = Inches(0.5)
    else: new_p.paragraph_format.first_line_indent = Inches(0)
    return new_p

# Isi Latar Belakang (setelah index 125)
add_after(doc, 125, "Komunikasi interpersonal merupakan proses pertukaran pesan antara dua individu atau lebih yang terjadi secara langsung atau tidak langsung, dengan tujuan saling memahami dan mempengaruhi satu sama lain. Dalam konteks keluarga, komunikasi interpersonal antara anak dan orang tua memiliki peran yang sangat penting dalam membentuk karakter, motivasi belajar, serta kesejahteraan psikologis anak (DeVito, 2016). Sejak pandemi COVID-19, sistem pembelajaran di Indonesia mengalami perubahan signifikan dengan diterapkannya pembelajaran jarak jauh atau Belajar dari Rumah (BDR). Perubahan ini tidak hanya mempengaruhi proses pembelajaran di sekolah, tetapi juga mengubah dinamika interaksi antara anak dan orang tua di rumah.", align='justify')
add_after(doc, 125, "Di Kabupaten Bantul, Daerah Istimewa Yogyakarta, banyak orang tua yang harus berperan ganda: sebagai pengasuh, pendamping belajar, sekaligus komunikator utama bagi anak-anak mereka selama masa BDR. Fenomena ini menimbulkan berbagai tantangan dalam pola komunikasi interpersonal, seperti perubahan frekuensi interaksi, pergeseran topik percakapan dari aktivitas sehari-hari menjadi fokus pada tugas belajar, serta munculnya konflik akibat tekanan akademik yang dirasakan bersama oleh anak dan orang tua. Oleh karena itu, penelitian ini bertujuan untuk memahami bagaimana pola komunikasi interpersonal anak dengan orang tua berlangsung dalam pendampingan BDR di Bantul.", align='justify')

print("Latar belakang ditambahkan")

# Isi Identifikasi Masalah (setelah index 126 - tapi karena kita sudah tambah 2 paragraf, index bergeser)
# Kita akan menggunakan pendekatan berbasis teks: cari paragraf dengan teks 'Identifikasi Masalah' lalu tambah setelahnya
for i, p in enumerate(doc.paragraphs):
    if p.text == "Identifikasi Masalah":
        add_after(doc, i, "Berdasarkan latar belakang di atas, beberapa masalah yang dapat diidentifikasi antara lain:", align='justify')
        items_id = [
            "Adanya perubahan signifikan dalam frekuensi dan kualitas komunikasi antara anak dan orang tua sejak diterapkannya BDR.",
            "Munculnya ketegangan atau konflik komunikasi akibat peran ganda orang tua sebagai pendamping belajar dan pengasuh.",
            "Kurangnya pemahaman orang tua terhadap strategi komunikasi yang efektif dalam mendukung proses belajar anak dari rumah.",
            "Terbatasnya penelitian yang secara spesifik mengkaji pola komunikasi interpersonal anak-orang tua dalam konteks BDR di wilayah Bantul."
        ]
        for item in items_id:
            add_after(doc, i + 1 if i == len(doc.paragraphs)-1 else doc.paragraphs.index(p)+1, "\u2022  " + item, align='justify', indent_first=True)
        # Sebenarnya ini akan berulang; mari sederhanakan
        break
print("Identifikasi masalah ditambahkan (cek)")

doc.save('Downloads/Proposal_Penelitian_Tia_Bab1.docx')
print("File disimpan sementara")
